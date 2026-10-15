import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from nexo_processamento.main import app

SEGREDO = "segredo-de-teste"
H = {"X-Nexo-Segredo": SEGREDO}


@pytest.fixture(autouse=True)
def _segredo(monkeypatch):
    monkeypatch.setenv("NEXO_PROCESSAMENTO_SEGREDO", SEGREDO)


@pytest.fixture
def cliente():
    return TestClient(app)


def pdf(paginas: int, com_texto: bool = True) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for i in range(1, paginas + 1):
        if com_texto:
            c.setFont("Helvetica-Bold", 20)
            c.drawString(60, 790, f"Capitulo {i}")
            c.setFont("Helvetica", 11)
            y = 760
            for par in range(3):
                for linha in range(4):
                    c.drawString(60, y, f"Pagina {i} paragrafo {par} linha {linha}: texto de estudo sobre o assunto.")
                    y -= 14
                y -= 14
        else:
            c.rect(100, 100, 200, 200)  # só desenho: simula PDF escaneado sem texto
        c.showPage()
    c.save()
    return buf.getvalue()


def docx(titulos: bool = True) -> bytes:
    d = Document()
    if titulos:
        d.add_heading("Introducao", level=1)
    d.add_paragraph("Primeiro paragrafo do documento com conteudo suficiente. " * 8)
    d.add_heading("Desenvolvimento", level=1)
    d.add_paragraph("Segundo paragrafo, ja na outra secao do documento. " * 8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
