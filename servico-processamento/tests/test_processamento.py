import io
import zipfile

from nexo_processamento import extracao
from nexo_processamento.divisao import dividir
from nexo_processamento.extracao import Bloco
from nexo_processamento.ia.linguagem import RespostaLinguagem
from nexo_processamento.topicos import cobre, sugerir
from conftest import H, docx, pdf


def _enviar(cliente, rota, dados, headers=H, nome="a.bin"):
    return cliente.post(rota, files={"arquivo": (nome, dados)}, headers=headers)


# --- autenticação entre serviços
def test_sem_segredo_ou_errado_recusa(cliente):
    assert _enviar(cliente, "/interno/inspecionar", b"x", headers={}).status_code == 401
    assert _enviar(cliente, "/interno/inspecionar", b"x", headers={"X-Nexo-Segredo": "errado"}).status_code == 401
    assert cliente.post("/interno/embeddings", json={"textos": ["a"]}).status_code == 401
    assert cliente.post("/interno/topicos", json={"trechos": [{"ordem": 1}]}).status_code == 401


def test_servico_sem_segredo_configurado_nao_aceita_nada(cliente, monkeypatch):
    monkeypatch.delenv("NEXO_PROCESSAMENTO_SEGREDO")
    assert _enviar(cliente, "/interno/inspecionar", b"x").status_code == 503


# --- inspeção e formato pelo conteúdo
def test_inspeciona_pdf_docx_txt(cliente):
    assert _enviar(cliente, "/interno/inspecionar", pdf(3)).json() == {"tipo": "pdf", "paginas": 3, "caracteres": None}
    assert _enviar(cliente, "/interno/inspecionar", docx()).json()["tipo"] == "docx"
    assert _enviar(cliente, "/interno/inspecionar", "olá".encode()).json() == {"tipo": "txt", "paginas": None, "caracteres": 3}


def test_binario_zip_comum_e_pdf_corrompido_sao_recusados(cliente):
    z = io.BytesIO()
    with zipfile.ZipFile(z, "w") as f:
        f.writestr("a.txt", "x")
    for dados in (b"\x00\x01\x02binario", z.getvalue(), b"%PDF-1.4 lixo sem estrutura", b"PK\x03\x04quebrado"):
        r = _enviar(cliente, "/interno/inspecionar", dados)
        assert r.status_code == 422 and r.json()["detail"]["codigo"] == "formato_invalido", dados[:10]


def test_docx_grande_demais_descompactado_e_recusado(monkeypatch):
    monkeypatch.setattr(extracao, "MAX_DESCOMPACTADO", 10)
    try:
        extracao.detectar(docx())
        assert False
    except extracao.FormatoInvalido:
        pass


# --- extração
def test_pdf_com_numero_de_pagina_e_titulos(cliente):
    r = _enviar(cliente, "/interno/processar", pdf(3)).json()
    assert r["paginas"] == 3 and r["tipo"] == "pdf"
    t = r["trechos"]
    assert [x["ordem"] for x in t] == list(range(1, len(t) + 1))
    assert {x["numero_pagina"] for x in t} == {1, 2, 3}
    assert [x["titulo"] for x in t if x["titulo"]] == ["Capitulo 1", "Capitulo 2", "Capitulo 3"]  # título abre trecho novo


def test_pdf_sem_texto_e_sem_texto_extraivel(cliente):
    r = _enviar(cliente, "/interno/processar", pdf(2, com_texto=False))
    assert r.status_code == 422 and r.json()["detail"]["codigo"] == "sem_texto"


def test_docx_usa_estilos_de_titulo(cliente):
    t = _enviar(cliente, "/interno/processar", docx()).json()["trechos"]
    assert [x["titulo"] for x in t] == ["Introducao", "Desenvolvimento"] and all(x["numero_pagina"] is None for x in t)


def test_txt_utf8_e_cp1252(cliente):
    corpo = "Primeiro parágrafo com acentuação.\n\nSegundo parágrafo, também com acentuação."
    for dados in (corpo.encode("utf-8"), corpo.encode("cp1252")):
        t = _enviar(cliente, "/interno/processar", dados).json()["trechos"]
        assert "acentuação" in t[0]["conteudo"] and "parágrafo" in t[0]["conteudo"]


def test_txt_vazio_e_sem_texto(cliente):
    assert _enviar(cliente, "/interno/processar", b"   \n\n  ").json()["detail"]["codigo"] == "sem_texto"


def test_arquivo_acima_do_teto(cliente, monkeypatch):
    monkeypatch.setattr("nexo_processamento.main.TAMANHO_MAXIMO", 10)
    assert _enviar(cliente, "/interno/processar", b"a" * 11).status_code == 413


# --- divisão em trechos
def test_divisao_respeita_tamanho_titulo_e_nao_sobrepoe():
    par = "palavra " * 60  # ~480 caracteres
    blocos = [Bloco(par, 1), Bloco(par, 1), Bloco(par, 1), Bloco(par, 1), Bloco("Novo capitulo", 2, True), Bloco(par, 2)]
    t = dividir(blocos, alvo=1000, maximo=1200)
    assert [x.ordem for x in t] == list(range(1, len(t) + 1))
    assert all(len(x.conteudo) <= 1200 + 10 for x in t)
    assert t[-1].titulo == "Novo capitulo" and t[-1].conteudo.startswith("Novo capitulo")
    assert sum(x.conteudo.count("palavra") for x in t) == 5 * 60  # nada perdido nem duplicado


def test_paragrafo_gigante_e_cortado_em_fim_de_frase_e_sem_frase():
    com_frases = "Uma frase com algumas palavras. " * 100
    t = dividir([Bloco(com_frases, 1)], alvo=500, maximo=600)
    assert all(len(x.conteudo) <= 600 for x in t) and all(x.conteudo.endswith(".") for x in t)
    t = dividir([Bloco("x" * 2500, 1)], alvo=500, maximo=600)
    assert all(len(x.conteudo) <= 600 for x in t) and sum(len(x.conteudo) for x in t) == 2500


# --- embeddings em lotes
def test_embeddings(cliente):
    r = cliente.post("/interno/embeddings", json={"textos": ["a", "b"]}, headers=H).json()
    assert len(r["vetores"]) == 2 and len(r["vetores"][0]) == 1536 and r["modelo"] == "embedding-falso"
    assert cliente.post("/interno/embeddings", json={"textos": ["a"] * 65}, headers=H).status_code == 422
    assert cliente.post("/interno/embeddings", json={"textos": []}, headers=H).status_code == 422


# --- sugestão de tópicos
class Roteirizado:
    nome = "roteirizado"

    def __init__(self, texto):
        self.texto = texto

    def gerar(self, prompt, imagens=None):
        return RespostaLinguagem(self.texto, 10, 5)


TRECHOS = [{"ordem": i, "titulo": None, "inicio": f"t{i}"} for i in range(1, 7)]


def test_cobertura():
    assert cobre([{"nome": "a", "inicio": 1, "fim": 3}, {"nome": "b", "inicio": 4, "fim": 6}], 6)
    for ruim in (
        [{"nome": "a", "inicio": 2, "fim": 6}],  # não começa em 1
        [{"nome": "a", "inicio": 1, "fim": 3}, {"nome": "b", "inicio": 5, "fim": 6}],  # lacuna
        [{"nome": "a", "inicio": 1, "fim": 4}, {"nome": "b", "inicio": 4, "fim": 6}],  # sobreposição
        [{"nome": "a", "inicio": 1, "fim": 5}],  # não termina em N
        [{"nome": " ", "inicio": 1, "fim": 6}],  # sem nome
        [],
    ):
        assert not cobre(ruim, 6), ruim


def test_sugestao_valida_aceita_e_invalida_vira_vazia():
    ok = '```json\n[{"nome":"Parte 1","inicio":1,"fim":2},{"nome":"Parte 2","inicio":3,"fim":6}]\n```'
    p = sugerir(Roteirizado(ok), TRECHOS)
    assert [t["nome"] for t in p.topicos] == ["Parte 1", "Parte 2"] and p.tokens_entrada == 10
    for lixo in ("não sei", '[{"nome":"x","inicio":1,"fim":3}]', '[{"nome":"x"}]', "[1,2]"):
        assert sugerir(Roteirizado(lixo), TRECHOS).topicos == []


def test_endpoint_topicos_com_modelo_falso_nao_da_proposta(cliente):
    r = cliente.post("/interno/topicos", json={"trechos": [{"ordem": 1, "inicio": "a"}]}, headers=H).json()
    assert r["topicos"] == [] and r["modelo"] == "linguagem-falso"


def test_pdf_de_50_paginas_processa(cliente):
    r = _enviar(cliente, "/interno/processar", pdf(50)).json()
    assert r["paginas"] == 50 and 50 <= len(r["trechos"]) <= 200


def test_txt_que_comeca_com_pk_continua_sendo_texto(cliente):
    dados = "PKI é a infraestrutura de chaves públicas; este é um texto comum.".encode()
    assert _enviar(cliente, "/interno/inspecionar", dados).json()["tipo"] == "txt"
    assert _enviar(cliente, "/interno/processar", dados).status_code == 200


def test_cliente_da_openai_e_reaproveitado(monkeypatch):
    from nexo_processamento.config import criar_embedding

    monkeypatch.setenv("EMBEDDING_PROVEDOR", "openai")
    monkeypatch.setenv("OPENAI_API_KEY", "k")
    assert criar_embedding() is criar_embedding()


def test_docx_com_tabela_entra_na_ordem_do_documento(cliente):
    from docx import Document

    d = Document()
    d.add_paragraph("Antes da tabela, um paragrafo introdutorio.")
    t = d.add_table(rows=2, cols=2)
    for i, r in enumerate(t.rows):
        for j, c in enumerate(r.cells):
            c.text = f"celula {i}{j}"
    t.cell(0, 0).merge(t.cell(0, 1))  # célula mesclada não repete o texto
    d.add_paragraph("Depois da tabela, outro paragrafo de conclusao.")
    buf = io.BytesIO()
    d.save(buf)
    texto = _enviar(cliente, "/interno/processar", buf.getvalue()).json()["trechos"][0]["conteudo"]
    assert texto.index("introdutorio") < texto.index("celula 00") < texto.index("conclusao")
    assert texto.count("celula 00") == 1 and "celula 10 | celula 11" in texto


def test_docx_so_com_tabela_nao_e_sem_texto(cliente):
    from docx import Document

    d = Document()
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "primeira coluna da tabela", "segunda coluna da tabela"
    buf = io.BytesIO()
    d.save(buf)
    assert _enviar(cliente, "/interno/processar", buf.getvalue()).status_code == 200


def test_secoes_curtas_se_juntam_em_vez_de_virar_trechos_minusculos():
    blocos = []
    for i in range(60):
        blocos += [Bloco(f"Titulo {i}", 1, True), Bloco("Texto curto da secao.", 1)]
    t = dividir(blocos)
    assert len(t) < 15 and all(len(x.conteudo) >= 100 for x in t[:-1])
    assert sum(x.conteudo.count("Texto curto") for x in t) == 60  # nada perdido


def test_titulo_depois_de_corpo_suficiente_ainda_abre_trecho_novo():
    corpo = "palavra " * 60  # ~480 caracteres, acima do mínimo
    t = dividir([Bloco(corpo, 1), Bloco("Capitulo 2", 2, True), Bloco(corpo, 2)])
    assert len(t) == 2 and t[1].conteudo.startswith("Capitulo 2") and t[1].titulo == "Capitulo 2"


def test_docx_niveis_de_titulo_e_capitulo_sempre_abre_trecho(cliente):
    from docx import Document

    d = Document()
    d.add_heading("Capitulo 1", level=1)
    d.add_paragraph("curto.")                   # corpo curto: antes, o capítulo seguinte se juntava a este trecho
    d.add_heading("Capitulo 2", level=1)
    d.add_paragraph("outro texto curto.")
    d.add_heading("Secao 2.1", level=2)
    d.add_paragraph("texto da subsecao.")
    buf = io.BytesIO()
    d.save(buf)
    t = _enviar(cliente, "/interno/processar", buf.getvalue()).json()["trechos"]
    assert [(x["titulo"], x["nivel"]) for x in t if x["nivel"]] == [("Capitulo 1", 1), ("Capitulo 2", 1)]  # cada capítulo abre o seu trecho
    assert "Secao 2.1" in t[-1]["conteudo"] and t[-1]["nivel"] == 1   # subseção curta entra no trecho do capítulo


def test_nivel_do_estilo():
    from nexo_processamento.extracao import _nivel_do_estilo as n

    assert [n("Heading 1"), n("Heading 3"), n("Título 2"), n("Title"), n("Normal"), n("List Bullet"), n(None)] == [1, 3, 2, 1, None, None, None]


def test_pdf_niveis_pela_ordem_do_tamanho_da_fonte(cliente):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFont("Helvetica-Bold", 24); c.drawString(60, 780, "Capitulo Grande")
    c.setFont("Helvetica-Bold", 16); c.drawString(60, 740, "Secao Media")
    c.setFont("Helvetica", 10)
    for i in range(10):
        c.drawString(60, 700 - i * 14, f"linha de corpo numero {i} com texto suficiente para ser o corpo")
    c.save()
    doc = extracao.extrair(buf.getvalue())
    assert [(b.texto, b.nivel) for b in doc.blocos if b.titulo] == [("Capitulo Grande", 1), ("Secao Media", 2)]


def test_prompt_de_topicos_leva_o_nivel_e_pede_para_seguir_a_estrutura():
    from nexo_processamento.topicos import _prompt

    p = _prompt([{"ordem": 1, "nivel": 1, "titulo": "5. Área 2: edição", "inicio": "texto"}, {"ordem": 2, "nivel": None, "titulo": None, "inicio": "mais"}])
    assert "1\t1\t5. Área 2: edição\ttexto" in p and "2\t\t\tmais" in p
    assert "Siga a estrutura" in p and "nível 1" in p


# --- reorganização com IA (botão da tela de revisão)
ATUAIS = [{"nome": "A", "inicio": 1, "fim": 4, "congelado": False}, {"nome": "B", "inicio": 5, "fim": 6, "congelado": True}]


def test_prompt_de_reorganizar_leva_atuais_congelados_e_estrutura():
    from nexo_processamento.topicos import _prompt_refinar

    p = _prompt_refinar([{"ordem": i, "nivel": 1 if i == 1 else None, "titulo": "Cap" if i == 1 else None, "inicio": "x"} for i in range(1, 7)], ATUAIS)
    assert "ATUAIS_JSON: " in p and '"congelado": true' in p and "congelado=true devem ser mantidos EXATAMENTE" in p
    assert "1\t1\tCap\tx" in p and "TODOS os trechos de 1 a 6" in p


def test_refinar_aceita_proposta_valida_que_preserva_congelados():
    from nexo_processamento.topicos import refinar

    r = refinar(Roteirizado('[{"nome":"Tudo antes","inicio":1,"fim":4},{"nome":"B renomeado","inicio":5,"fim":6}]'), TRECHOS, ATUAIS)
    assert [t["nome"] for t in r.topicos] == ["Tudo antes", "B renomeado"] and r.tokens_entrada == 10


def test_refinar_recusa_cobertura_ruim_e_tocar_em_congelado():
    from nexo_processamento.topicos import refinar

    for ruim in (
        '[{"nome":"A","inicio":1,"fim":4}]',                                          # não cobre até o fim
        '[{"nome":"A","inicio":1,"fim":3},{"nome":"B","inicio":4,"fim":6}]',          # mexeu no congelado (5-6)
        "sem json",
    ):
        assert refinar(Roteirizado(ruim), TRECHOS, ATUAIS).topicos == [], ruim


def test_modelo_falso_devolve_a_divisao_atual_na_reorganizacao(cliente):
    corpo = {"trechos": [{"ordem": i, "inicio": "t"} for i in range(1, 7)], "atuais": ATUAIS}
    r = cliente.post("/interno/reorganizar", json=corpo, headers=H).json()
    assert [(t["inicio"], t["fim"]) for t in r["topicos"]] == [(1, 4), (5, 6)] and r["modelo"] == "linguagem-falso"


def test_reorganizar_exige_o_segredo_interno(cliente):
    assert cliente.post("/interno/reorganizar", json={"trechos": [{"ordem": 1}], "atuais": ATUAIS}).status_code == 401
